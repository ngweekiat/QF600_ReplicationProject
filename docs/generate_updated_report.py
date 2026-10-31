"""Build the compact report with Chapters 1–4 drafted and later guidance in red."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from jkmp.config import load_config

ROOT = Path(__file__).resolve().parent
cfg = load_config()
source = Document(ROOT / 'QF600_Replication_Framework.docx')
doc = Document()
BLACK, RED = RGBColor(0, 0, 0), RGBColor(192, 0, 0)
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Inches(0.7)
sec.left_margin = sec.right_margin = Inches(0.75)
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
style = doc.styles['Normal']
style.font.name, style.font.size, style.font.color.rgb = 'Calibri', Pt(10.5), BLACK
style.paragraph_format.space_after = Pt(6)
style.paragraph_format.line_spacing = 1.05
for name in ('Title', 'Subtitle', 'Heading 1', 'Heading 2'):
    doc.styles[name].font.color.rgb = BLACK
    doc.styles[name].paragraph_format.keep_with_next = True
doc.styles['Heading 1'].font.size = Pt(16)
doc.styles['Heading 2'].font.size = Pt(12)

def paragraph(text, red=False):
    p = doc.add_paragraph(text)
    for run in p.runs:
        run.font.color.rgb = RED if red else BLACK
    return p

def heading(text, level=1, red=False):
    p = doc.add_heading(text, level)
    for run in p.runs:
        run.font.color.rgb = RED if red else BLACK
    return p

def copy(prefix):
    matches = [p.text for p in source.paragraphs if p.text.startswith(prefix)]
    assert len(matches) == 1, prefix
    return paragraph(matches[0])

def equation(text):
    p = paragraph(text)
    p.paragraph_format.left_indent = Inches(0.15)
    for run in p.runs:
        run.font.name = 'Cambria'

def table(headers, rows, red=False, compact=False):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Light Shading Accent 1'
    for cell, text in zip(t.rows[0].cells, headers):
        cell.text = str(text)
    for row in rows:
        for cell, text in zip(t.add_row().cells, row):
            cell.text = str(text)
    for i, row in enumerate(t.rows):
        props = row._tr.get_or_add_trPr()
        props.append(OxmlElement('w:cantSplit'))
        if i == 0:
            props.append(OxmlElement('w:tblHeader'))
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(3)
                for run in p.runs:
                    run.font.color.rgb = RED if red else BLACK
                    run.font.size = Pt(8.5 if compact else 9.5)
    return t

heading('Static-ML* versus Portfolio-ML', 0)
paragraph('QF600 Asset Pricing | Modified Python replication | 11 October 2026')
paragraph('Draft for a report limited to 20 pages, including references. No separate title page is used.')
p = paragraph('Colour guide: Black = implemented methods, settings and verified status. ')
p.add_run('Red = hypotheses, pending-result instructions and guidance for further work.').font.color.rgb = RED

heading('1. Executive Summary')
copy('This project implements a modified Python replication')
copy('The current configuration targets approximately')
paragraph('The central hypothesis is that directly learning a portfolio improves net investor utility relative to predicting returns first and then optimising. Portfolio-ML may benefit from aligning estimation with the investor’s economic objective, but superiority must be established using the replication’s out-of-sample results.', red=True)
paragraph('Empirical results are not yet available because the WRDS authentication and licensed data download remain incomplete. The current document therefore explains the implemented research design and provides instructions for completing the results and extension chapters. Software checks passed 132 tests with three skips; those checks are not evidence of investment performance.')

doc.add_page_break()
heading('2. Introduction: Research Problem and Business Motivation')
paragraph('The research question is whether Portfolio-ML offers higher net investor utility than Static-ML* when both use the same information, eligible stocks, covariance estimates, trading-cost assumptions and investor settings. The business problem is the construction of a scalable portfolio: useful forecasts must translate into positions whose returns justify their risk and execution costs.')
paragraph('Static-ML* is a meaningful comparator because it already considers trading costs and existing holdings when choosing weights. The experiment therefore examines whether learning the portfolio directly adds value beyond a cost-aware predict-then-optimise procedure. It does not compare Portfolio-ML only against a portfolio that ignores implementation costs.')
paragraph('A persistent, liquid signal may have greater investment value than a more accurate but expensive-to-trade signal. Portfolio-ML may exploit this difference because its learning objective incorporates risk and trading costs. This potential mechanism motivates the comparison; its practical importance remains an empirical question.', red=True)
paragraph('The project is a modified replication of Jensen, Kelly, Malamud and Pedersen’s framework [1], using the authors’ R repository as a reference [2]. Changes include a smaller, buffered NYSE/Nasdaq universe, 50 selected characteristics, a later test period, a broader top-500 covariance-estimation universe and a focused two-method comparison. These changes mean the original paper’s numerical results should not be expected to reproduce exactly.')
paragraph('The report proceeds from the dataset and portfolio rules to out-of-sample evaluation, then examines three controlled iterations and macroeconomic stress conditions. The final chapter translates the evidence into implications for investor size, trading discipline and model selection.')
paragraph('Page allocation: Executive Summary 1; Introduction 1; Data and Methodology 4; Performance and Results 4; Iterations and Assumptions 4; Macro Stress Tests 3; Conclusion 1; References 2. Total ceiling: 20 pages. Replace guidance with final analysis rather than adding it alongside the completed results.', red=True)

doc.add_page_break()
heading('3. Data and Methodology')
heading('3.1 Data Sources, Sample Periods, and Stock Universe', 2)
paragraph('The configured monthly source is JKP Global Factor Data on WRDS (contrib.global_factor). Daily returns are configured to use CRSP’s CIZ file (crsp.dsf_v2). Public inputs include JKP Factor Details and Cluster Labels plus Ken French monthly and daily factors. Source-table availability, column definitions and sample coverage must be checked on the actual account when the licensed download is resumed.')
table(['Component', 'Implemented specification'], [
    ('Data period', 'Warm-up starts January 1993; test returns January 2020–December 2025'),
    ('Model date labels', 'Training 2005–2014; initial validation 2015–2019; annual expanding refits and selections thereafter'),
    ('Portfolio universe', 'Target 80 eligible US stocks; NYSE and Nasdaq; market-capitalisation selection'),
    ('Membership', 'December reconstitution; entry rank ≤70; incumbent retention rank ≤100; fill remaining slots to target'),
    ('Screens', 'Share codes 10/11/12; JKP common/primary_sec/obs_main flags; positive market cap and volume; valid returns, industry data and consecutive history'),
    ('Investor', 'Risk aversion γ=10; wealth $10B at end-2025; common market-linked wealth path'),
])
copy('The 80-stock target is not a strict monthly count')
paragraph('The warm-up period provides a full 2,520-day covariance window before model evaluation. Portfolio selection and screening use current and historical information. This reduces mechanical forward-looking selection, although point-in-time data availability and revision conventions still need to be verified against the supplied database fields.')

heading('3.2 Characteristic Selection and Data Processing', 2)
paragraph('Both methods use 50 underlying stock characteristics across 13 JKP themes. The set was selected in advance for economic coverage from the original R code’s retained features, and checked against the downloaded JKP characteristic definitions and theme labels. It is not a ranking of the 50 most predictive characteristics. The code requires at least 25 of the 50 inputs to be available, then applies monthly cross-sectional percentile ranks and imputes missing ranks to 0.5.')
table(['Theme', 'Count', 'Selected JKP characteristics'], [
    (theme.replace('_', ' ').title(), len(names), ', '.join(names))
    for theme, names in cfg.features.themes.items()
], compact=True)
paragraph('Some theme assignments are not intuitive from the characteristic name: for example, the JKP size cluster includes liquidity measures, and profit growth includes ret_12_7. The replication follows the metadata rather than redefining these clusters. Random Fourier features are transformations of the 50 inputs, not additional observed stock characteristics.')
paragraph('Portfolio-ML uses current and lagged characteristic signals, with an additional historical observation needed for changes in the lagged signal portfolio. Stocks leaving eligibility during a membership year cannot re-enter before the next annual reconstitution. Increasing the feature set from 18 to 50 also changes coverage screens and covariance theme factors, so earlier prepared data and fitted models must be rebuilt.')

heading('3.3 Implementable Efficient Frontier: Trading Costs and Investor Size', 2)
paragraph('Jensen et al. define the investment opportunity set in terms of expected returns net of trading costs at each risk level [1]. The local implementation operationalises this through a common covariance model, a quadratic price-impact specification and portfolios evaluated after trading costs. Investor wealth affects the cost of changing portfolio weights.')
paragraph('The covariance model uses daily returns from a broader top-500 market-capitalisation universe and combines six sector groups with 13 characteristic themes. It requires a full ten-year daily window, applies exponentially weighted covariance estimation, and scales daily covariance to monthly covariance using 21 trading days. Idiosyncratic variances have minimum-observation requirements and fallback rules.')
equation('TCₜ = ½ Wₜ (wₜ − wₜˢᵗᵃʳᵗ)ᵀ Λₜ (wₜ − wₜˢᵗᵃʳᵗ).')
equation('Λᵢ,ₜ = 2 × price_impact / dollar_volumeᵢ,ₜ = 0.2 / dollar_volumeᵢ,ₜ.')
paragraph('Here wₜˢᵗᵃʳᵗ denotes inherited weights after return drift, Wₜ is investor wealth and Λₜ contains stock-level impact coefficients. Dollar volume uses the 126-day measure. The current configuration assumes dollar units, which must be verified after downloading the actual data. The model explicitly charges market impact but does not separately model stock-borrow fees or financing spreads.')
paragraph('A larger investor incurs higher cost for the same proportional weight change. The common wealth path is rolled backwards from end-2025 wealth using market total returns; it is not each strategy’s compounded realised wealth. Sales of departing stocks incur costs when trading data remain available. The frontier analysis retunes both methods over γ∈{1,5,10,20,100} and terminal wealth {$1B,$10B,$50B,$100B,$250B}.')

heading('3.4 Static-ML* versus Portfolio-ML: Estimation and Portfolio Construction', 2)
paragraph('Static-ML* first predicts next-month excess returns using random Fourier features and ridge regression. The return model selects its settings by validation mean-squared error and refits before producing the following year’s predictions. For each prediction year, it uses a five-year historical validation window and an earlier expanding training sample. The configured feature-count grid is {2,4,8,16,32,64,128,256,512}.')
equation('wₜ = [γΣₜ + g diag(γΣₜ) + WₜkΛₜ]⁻¹ [u μ̂ₜ + WₜkΛₜ wₜˢᵗᵃʳᵗ].')
paragraph('The portfolio rule uses covariance Σₜ, predicted excess returns μ̂ₜ, trading-cost coefficients and existing holdings. Static-ML* tunes u∈{0.25,0.5,1}, g∈{0,1,2} and k∈{1,1/3,1/5}, representing return shrinkage, variance inflation and cost scaling. Each December, cumulative realised utility determines the adjustments for the next return year. “Static” refers to a one-period objective; the portfolio still rebalances monthly.')
paragraph('Portfolio-ML estimates coefficients for portfolio signals using an objective that integrates realised returns, covariance risk and trading costs. Its nonlinear feature-count grid is {64,128,256,512}. Signals are cross-sectionally standardised, scaled by stock volatility and combined across historical observations using trading-speed and return-drift matrices. Coefficients are estimated using only prior realised-return observations, and settings are selected annually by cumulative realised utility.')
equation('β̂ = (Σ̄̃ + λI)⁻¹ r̄̃;     wₜ = mₜwₜˢᵗᵃʳᵗ + (I − mₜ)aimₜ.')
paragraph('The transformed covariance moment Σ̄̃ combines risk and trading-cost terms. The matrix mₜ controls retention of inherited holdings, so I−mₜ controls adjustment towards the learned aim. The code computes mₜ with ten fixed-point iterations. “Aiming in front of the target” describes anticipation under costly adjustment; the implementation uses current and lagged signals, not knowledge of future realised returns or separate forecasts for every future horizon.')
paragraph('Ridge regularisation controls coefficient instability and is distinct from investor risk aversion. The shared grid contains zero and 100 exponentially spaced penalties from exp(−10) to exp(10). Portfolio-ML penalises its constant signal along with the other coefficients. Annual choices apply to subsequent return years rather than being selected using the same year’s test returns.')
paragraph('Portfolio-ML may outperform because it learns which signals are valuable after risk and implementation costs, rather than selecting them solely for stock-level forecasting accuracy. This alignment is a plausible advantage, not a guaranteed outcome: Static-ML* already accounts for costs, and finite-sample estimation error can affect either method.', red=True)
paragraph('The implementation does not feed Markowitz-ML weights into Static-ML*. The active workflow uses preparation, covariance estimation and return prediction, skips the benchmark stage, then fits Static-ML* and Portfolio-ML. Older benchmark and linear-extension results are excluded when rebuilding the selected comparison.')

doc.add_page_break()
heading('4. Performance Evaluation and Replication Results')
heading('4.1 Evaluation Metrics', 2)
paragraph('Weights formed at month-end t earn returns over t+1. Both test portfolios start from value-weighted holdings and then inherit return-drifted positions. Let r be monthly gross excess return, TC monthly trading cost as a fraction of wealth, and s(r) the sample standard deviation. The implementation uses n−1 sample variance.')
table(['Metric', 'Current calculation'], [
    ('Annual gross excess return', '12 × mean(r)'),
    ('Annual net excess return', '12 × mean(r−TC)'),
    ('Annual volatility', '√12 × s(r)'),
    ('Gross Sharpe ratio', '√12 × mean(r)/s(r)'),
    ('Net Sharpe ratio', '√12 × mean(r−TC)/s(r)'),
    ('Annual trading costs', '12 × mean(TC)'),
    ('Annual investor utility', '12 × [mean(r)−mean(TC)−½γ Var(r)]'),
    ('Monthly turnover', 'Mean of Σ|w−wˢᵗᵃʳᵗ|; absolute buys and sells, without halving'),
    ('Gross and short exposure', 'Mean Σ|w| and mean Σ|max(−w,0)|'),
])
paragraph('The net Sharpe ratio uses gross-return volatility in the denominator, following the code’s convention. Annual returns and costs use arithmetic annualisation rather than compounded CAGR. Utility is a realised mean–variance criterion, not a statistical significance test. Turnover and exposures are monthly averages, not annualised.')

heading('4.2 Out-of-Sample Comparison: Static-ML* versus Portfolio-ML', 2)
paragraph('The selected test window contains returns from January 2020 through December 2025, giving 72 monthly observations if all required data are available. Annual refits and hyperparameter choices use historical observations available before the year being evaluated. Test-period observations can become historical inputs for later annual updates; this is expanding out-of-sample evaluation, not a single model frozen at December 2019.')
paragraph('The result-writing stage builds outputs/table2_test.csv with the two selected methods. Monthly returns, costs, holdings and annual choices are retained for verification. The common universe, investor assumptions and accounting rules support a like-for-like comparison. At present, the licensed input data and test portfolios have not been completed, so the actual sample count and performance differences are not established.')
table(['Required result', 'Static-ML*', 'Portfolio-ML'], [
    ('Net return / volatility / net Sharpe', 'Pending empirical run', 'Pending empirical run'),
    ('Investor utility / annual trading costs', 'Pending empirical run', 'Pending empirical run'),
    ('Turnover / gross exposure / short exposure', 'Pending empirical run', 'Pending empirical run'),
], red=True)
paragraph('Replace the pending cells with the actual result table. Report ΔU=U(Portfolio-ML)−U(Static-ML*) and explain whether its sign is supported by net Sharpe and trading-cost differences. Examine exposures before attributing higher returns to learning quality. State explicitly if the methods trade off return, risk and costs rather than one dominating the other.', red=True)

heading('4.3 Implementable Efficient Frontiers and Investor Capacity', 2)
paragraph('The frontier stage reruns and retunes the portfolios across 25 risk-aversion/wealth combinations. Its outputs summarise net return, volatility, net Sharpe, utility, costs and exposure. The plots include the two-method frontier at baseline wealth, the Portfolio-ML frontier by investor wealth, and a utility-versus-wealth capacity comparison. No Markowitz-ML line is required.')
paragraph('Insert the baseline frontier and capacity chart after running the grid. Discuss whether Portfolio-ML produces higher net return at comparable risk and whether its utility advantage survives increasing investor size. Distinguish an advantage at selected grid points from dominance over the full opportunity set. The origin drawn in a frontier plot is a modelling reference, not a separately estimated observation.', red=True)

heading('4.4 Economic Feature Importance', 2)
paragraph('Feature importance is evaluated through theme-level permutation across stocks within each month. The code preserves unshuffled characteristics and reuses the same monthly donor permutation across themes and methods. Portfolio-ML retains fitted coefficients while recomputing aims; Static-ML* recomputes predicted returns using its fitted yearly models, retains annual portfolio adjustments, and reruns its portfolio. Both are scored net of trading costs.')
equation('Importance(theme) = Utility(unshuffled) − Utility(theme shuffled).')
paragraph('A positive difference means that disrupting the theme reduced realised utility. A negative difference means the shuffled portfolio performed better in that particular draw. This measures dependence of the fitted strategy on the theme; it is not a causal effect and cannot uniquely allocate shared information between correlated themes. The default implementation uses one monthly permutation from seed 1 and reports point estimates.')
paragraph('Insert the theme-importance chart and identify the themes that contribute most to each method’s net utility. Do not infer that a theme is important solely because it contains more characteristics. Check alternative permutation seeds before describing a ranking as stable. Discuss differing theme use as a possible explanation for the performance comparison, rather than treating it as proof of superiority.', red=True)

doc.add_page_break()
heading('5. Iterative Improvements and Assumption Testing', red=True)
paragraph('Writing guidance: keep this chapter within four pages. Use a compact baseline-versus-iteration table and explain each change through motivation, modification, measured result and implication. The three designs below are proposed experiments, not completed improvements. Change one major assumption at a time and retain the baseline comparison.', red=True)
heading('5.1 Iteration 1: Motivation, Modification, and Results', 2, red=True)
paragraph('Suggested focus: trading-cost calibration. Verify dollar-volume units first. Then compare the baseline price-impact parameter of 0.1 with pre-specified lower and higher alternatives, such as 0.05 and 0.20. Rebuild affected inputs and retune both methods under each assumption. Report net utility, net Sharpe, costs and turnover in one table. Explain whether Portfolio-ML’s relative performance is sensitive to the assumed impact intensity.', red=True)
heading('5.2 Iteration 2: Motivation, Modification, and Results', 2, red=True)
paragraph('Suggested focus: universe breadth. Compare target universes of 50, 80 and 100 stocks while maintaining valid entry/retention buffers and identical dates. Rebuild membership and all dependent estimates. Report realised stock counts, coverage, concentration and performance. A larger universe might improve diversification or introduce less-liquid stocks; let the results establish which effect matters. These universe sizes are proposals, not changes already applied to the baseline.', red=True)
heading('5.3 Iteration 3: Motivation, Modification, and Results', 2, red=True)
paragraph('Suggested focus: model complexity and regularisation. Compare a pre-specified reduced random-feature grid with the baseline grid, and repeat the comparison across a small declared set of seeds. Keep the 50 underlying characteristics and the economic assumptions fixed. Examine whether a performance advantage survives simpler models and random-feature draws. Do not choose the final specification solely because it wins on 2020–2025 test outcomes.', red=True)
heading('5.4 Robustness and Key Limitations', 2, red=True)
paragraph('Summarise data coverage, revisions and timing assumptions; quadratic price impact; absence of explicit borrowing/financing fees; the small large-cap universe; the 72-month test period; and the gross-volatility denominator used for net Sharpe. Distinguish robustness checks actually run from recommended future work. If adding paired bootstrap uncertainty for ΔU or net Sharpe differences, state the resampling design and preserve dependence between the two methods. Such confidence intervals are not produced automatically by the current scripts.', red=True)

doc.add_page_break()
heading('6. Real-World Relevance: Macro Stress Tests', red=True)
paragraph('Writing guidance: reserve three pages. Reuse the Chapter 4 test portfolios instead of presenting the full-sample comparison again. The regime script already supports FRED recession, VIX and interest-rate labels plus VIX-based re-costing; results remain pending. Distinguish retrospective diagnostics from information that would have been available for live portfolio decisions.', red=True)
heading('6.1 Crisis Event Windows', 2, red=True)
paragraph('Use the configured descriptive windows: COVID crash (February–April 2020), 2022 hikes/bear market (January–December 2022), SVB/regional banks (March 2023), and the configured tariff-shock window (April 2025). Validate and cite the event labels before final submission. The current script shades these windows on the cumulative-utility chart; a dedicated event-level performance table requires additional aggregation. Report observation counts and avoid interpreting one-month annualised statistics as reliable estimates of persistent performance.', red=True)
heading('6.2 Crisis Re-Costing with Time-Varying Trading Costs', 2, red=True)
paragraph('Use the implemented diagnostic Λₜ(stress)=Λₜ×VIXₜ/mean(VIX), with the mean taken over the test window under the current configuration. Explain that this re-costs the existing trades: it does not retrain or re-optimise portfolios in response to the new costs. Compare stressed utility and cost changes for both methods. The full-window normalisation is an ex-post diagnostic; a live adaptive strategy would require a lagged or otherwise point-in-time scale and another portfolio run.', red=True)
heading('6.3 Regime-Conditional Performance', 2, red=True)
paragraph('Use USREC for recession/non-recession, test-window VIX terciles for volatility states, and three-month FEDFUNDS changes for hiking/cutting/on-hold states. The configured interest-rate threshold is ±0.10 percentage points. Show a compact utility-and-cost heatmap or table with sample counts. Explain that full-test-window VIX terciles are descriptive classifications, and that small or overlapping regime samples weaken broad claims about robustness. Do not label this a macro-timing strategy: the baseline portfolios are not trained to trade directly on these regime labels.', red=True)

doc.add_page_break()
heading('7. Conclusion and Business Implications', red=True)
paragraph('Writing guidance: keep the conclusion to one page. Answer the research question using the actual ΔU, net Sharpe comparison and frontier evidence. State whether Portfolio-ML improved the baseline outcome, under which investor-size and cost assumptions, and how stable the result was across the three iterations. If the evidence is mixed or favours Static-ML*, report that clearly.', red=True)
paragraph('Translate the evidence into the business problem: when might a larger investor benefit from direct portfolio learning, and when might a simpler predict-then-optimise approach remain competitive? Discuss liquidity, trade persistence, estimation stability and operational complexity. Tie every recommendation to a reported result rather than to the theoretical appeal of the method.', red=True)
paragraph('End with the main limitations and one or two evidence-based next steps. Do not describe hypothetical improvements as implemented, infer causality from permutation importance, or promise live performance from this historical exercise. Close with a concise assessment of what this modified replication establishes and what remains uncertain.', red=True)

doc.add_page_break()
heading('References')
paragraph('[1] Jensen, T. I., Kelly, B., Malamud, S., and Pedersen, L. H. (2026). Machine Learning and the Implementable Efficient Frontier. The Review of Financial Studies, 39(10), 3035–3078. https://doi.org/10.1093/rfs/hhag022.')
paragraph('[2] Jensen et al. Original replication code: https://github.com/theisij/ml-and-the-implementable-efficient-frontier. Local reference copies: reference/jkmp_r/.')
paragraph('[3] Local implementation: configs/base.yaml; src/jkmp/data/prepare.py and universe.py; src/jkmp/features.py and covariance.py; src/jkmp/models/return_model.py; src/jkmp/portfolio/static_ml.py, portfolio_ml.py, trading_speed.py and accounting.py; src/jkmp/evaluation/metrics.py, feature_importance.py and regimes.py; scripts/05_static_ml.py through 10_regimes.py; docs/Selected comparison.md and deviations.md.')
paragraph('Complete the final reference list with the exact JKP data release, CRSP documentation, Ken French factor definitions, FRED series documentation and sources for crisis-event dates. Include access dates and dataset vintages when known. Replace these instructions with final citations before submission.', red=True)

footer = sec.footer.paragraphs[0]
footer.alignment = 2
footer.add_run('QF600 replication | ')
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
doc.core_properties.title = 'QF600 Replication Report: Chapters 1–4 and Guidance'
doc.core_properties.author = 'QF600 Replication Project'
output = ROOT / 'QF600_Updated_Report_Draft.docx'
doc.save(output)

check = Document(output)
headings = [p.text for p in check.paragraphs if p.style.name.startswith('Heading')]
assert all(any(h.startswith(f'{i}. ') for h in headings) for i in range(1, 8))
assert len(cfg.features.names) == 50
later = False
for p in check.paragraphs:
    if p.text.startswith('5. Iterative'):
        later = True
    if p.text == 'References':
        later = False
    if later:
        assert all(r.font.color.rgb == RED for r in p.runs if r.text.strip())
words = sum(len(p.text.split()) for p in check.paragraphs)
words += sum(len(c.text.split()) for t in check.tables for row in t.rows for c in row.cells)
print(output)
print('Words:', words, '| Tables:', len(check.tables), '| Headings:', len(headings))
