"""Colour implemented descriptions black and hypotheses/proposed analysis red."""
from pathlib import Path
from docx import Document
from docx.shared import RGBColor

BLACK = RGBColor(0, 0, 0)
RED = RGBColor(192, 0, 0)
path = Path(__file__).resolve().parent / 'QF600_Replication_Framework.docx'
doc = Document(path)

# Exact passages, rather than keywords, keep implemented methods distinct from
# potential outcomes and suggested work within the same paragraph.
red_passages = [
    'The research hypothesis is that directly learning an investable portfolio can improve net investor utility relative to a two-stage procedure.',
    'A statistically useful signal may require frequent or expensive trades. Conversely, a persistent signal with modest forecasting accuracy may contribute more to a scalable investment strategy. The replication asks whether integrating implementation considerations into learning produces a better investment decision.',
    'Portfolio-ML must therefore demonstrate an advantage beyond simply subtracting costs from an otherwise cost-blind strategy.',
    'The key potential advantage is alignment: Portfolio-ML can prefer a signal because of its contribution to an implementable portfolio, even when it would not be preferred by a stock-level prediction criterion. However, Static-ML* also makes economic adjustments, and finite-sample estimation error can reverse the expected ranking.',
    'To show that Portfolio-ML is better in this replication, the empirical results must show higher out-of-sample net utility or a superior implementable frontier under matched assumptions. Higher gross return or lower prediction error alone is insufficient.',
    'The 72-month test window and single permutation draw can make rankings unstable, so sensitivity to alternative seeds and specifications is a useful subsequent analysis.',
    'These supplementary checks help identify whether an advantage is broad or concentrated in particular conditions.',
    'A credible conclusion should compare the utility difference U(Portfolio-ML)−U(Static-ML*), net Sharpe ratios, frontier positions and trading costs, while checking whether leverage or short exposure explains the ranking.',
    'paired uncertainty analysis would be an additional step.',
]

for style in doc.styles:
    if style.type == 1:
        style.font.color.rgb = BLACK

def black(paragraph):
    for run in paragraph.runs:
        run.font.color.rgb = BLACK

for paragraph in doc.paragraphs:
    black(paragraph)
    text = paragraph.text
    spans = []
    for passage in red_passages:
        start = text.find(passage)
        if start >= 0:
            spans.append((start, start + len(passage)))
    if not spans:
        continue
    paragraph.clear()
    offset = 0
    for start, end in sorted(spans):
        if start > offset:
            paragraph.add_run(text[offset:start]).font.color.rgb = BLACK
        paragraph.add_run(text[start:end]).font.color.rgb = RED
        offset = end
    if offset < len(text):
        paragraph.add_run(text[offset:]).font.color.rgb = BLACK

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                black(paragraph)
for section in doc.sections:
    for paragraph in list(section.header.paragraphs) + list(section.footer.paragraphs):
        black(paragraph)

if not any(p.text.startswith('Colour guide:') for p in doc.paragraphs):
    legend = doc.paragraphs[4].insert_paragraph_before()
    legend.add_run('Colour guide: ').bold = True
    legend.add_run('Black = implemented methods, current settings and verified status. ').font.color.rgb = BLACK
    legend.add_run('Red = hypotheses, potential advantages and proposed further analysis.').font.color.rgb = RED

try:
    doc.save(path)
except PermissionError:
    path = path.with_name('QF600_Replication_Framework_Coloured.docx')
    doc.save(path)

check = Document(path)
reds = [r.text for p in check.paragraphs for r in p.runs if r.font.color.rgb == RED]
assert len(reds) == len(red_passages) + 1, 'Some expected passages were not coloured.'
assert len(check.tables) == 2
print(path)
print('Red passages:', len(reds), '| Tables:', len(check.tables))
